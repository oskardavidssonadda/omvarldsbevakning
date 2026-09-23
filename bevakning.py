"""All logik för bevakningen. Appen (app.py) anropar funktionerna här."""
import calendar
import html
import re
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime, timezone
from io import BytesIO
from urllib.parse import urlparse
from zoneinfo import ZoneInfo

import feedparser
import pandas as pd
import requests
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import load_workbook

TIDSZON = ZoneInfo("Europe/Stockholm")
ORDFALT = ["nyckelord", "tvetydiga_ord", "kontextord", "kategoriord"]


def antal(n, en, flera):
    return f"{n} {en if n == 1 else flera}"


def vard(url):
    try:
        return urlparse(url).hostname.removeprefix("www.")
    except Exception:
        return url


# ---------- Bevakningar ----------

def tom_bevakning(namn="Min bevakning"):
    return {"namn": namn, "kallor": [], **{f: [] for f in ORDFALT}}


def normalisera(o):
    """Tar emot både nya filer och filer från den första versionen av appen."""
    b = tom_bevakning(o.get("namn") or o.get("rubrik") or "Min bevakning")
    if isinstance(o.get("kallor"), list):
        for k in o["kallor"]:
            if isinstance(k, (list, tuple)):
                namn, url, filtrera = (list(k) + [True])[:3]
            else:
                namn, url, filtrera = k.get("namn"), k.get("url"), k.get("filtrera", True)
            if url and not str(url).lower().startswith(("http://", "https://")):
                url = "https://" + str(url)
            if url:
                b["kallor"].append({"namn": namn or vard(url), "url": url, "filtrera": bool(filtrera)})
    else:
        for namn, url in (o.get("rena_kallor") or {}).items():
            b["kallor"].append({"namn": namn, "url": url, "filtrera": False})
        for namn, url in (o.get("blandade_kallor") or {}).items():
            b["kallor"].append({"namn": namn, "url": url, "filtrera": True})
    for f in ORDFALT:
        b[f] = rensa_ord(o.get(f) or [])
    return b


def rensa_ord(lista):
    """Städar en ordlista. Delar också upp inklistrade listor med komma, semikolon eller radbrytning."""
    ut = []
    for rad in lista:
        for o in re.split(r"[,;\n]", str(rad)):
            o = re.sub(r"^(?:[•\-*]\s+|\d+[.)]\s+)", "", o.strip()).strip(" \"'").lower()
            if o and o not in ut:
                ut.append(o)
    return ut


def las_kallor_fran_text(text):
    """Hittar källor i inklistrad text, t.ex. från en AI. En källa per rad: Namn | https://..."""
    kallor = []
    for rad in text.splitlines():
        m = re.search(r"https?://[^\s|<>\"')\]]+", rad)
        if not m:
            continue
        namn = re.sub(r"^(?:[•\-*]\s*|\d+[.)]\s*)", "", rad[:m.start()]).strip(" |:–-\t*[]()")
        url = m.group(0).rstrip(".,;")
        kallor.append({"namn": namn or vard(url), "url": url})
    return kallor


# ---------- Matchning ----------
# Ett ord matchar bara som helt ord. Med * i slutet matchar alla ord som börjar så,
# t.ex. ränt* -> räntan, räntor, räntehöjning.

def ord_monster(o):
    stjarna = o.endswith("*")
    bas = re.escape(o.rstrip("*"))
    return rf"(?<![^\W_]){bas}" + (r"[^\W_]*" if stjarna else r"(?![^\W_])")


def bygg_matchare(b):
    nyck = [(o, ord_monster(o)) for o in b["nyckelord"]]
    tvet = [(o, ord_monster(o)) for o in b["tvetydiga_ord"]]
    nyck_re = [(o, m, re.compile(m, re.I)) for o, m in nyck]
    tvet_re = [(o, m, re.compile(m, re.I)) for o, m in tvet]
    kont = [(k, ord_monster(k), re.compile(ord_monster(k), re.I)) for k in b["kontextord"] if k.rstrip("*")]
    kat = [k.rstrip("*") for k in b["kategoriord"] if k.rstrip("*")]

    def matcha(titel, sammanfattning, kategorier):
        text = f"{titel} {sammanfattning}".lower()
        for k in (k.lower() for k in kategorier):
            for o in kat:
                if o in k:
                    return f"kategori: {o}", []
        for o, m, r in nyck_re:
            if r.search(text):
                return o, [m]
        for o, m, r in tvet_re:
            if r.search(text):
                for k, km, kr in kont:
                    if kr.search(text):
                        return f"{o} + {k}", [m, km]
        return None, None

    return matcha


# ---------- Hämta ----------

def _rensa_html(text):
    return re.sub(r"\s+", " ", html.unescape(re.sub(r"<[^>]+>", " ", text or ""))).strip()


def _hamta_en(kalla):
    try:
        svar = requests.get(kalla["url"], timeout=12, headers={
            "User-Agent": "Mozilla/5.0 (compatible; Omvarldsbevakning/1.0)",
            "Accept": "application/rss+xml, application/atom+xml, application/xml, text/xml, */*"})
    except requests.Timeout:
        return kalla, None, "Källan svarade inte i tid"
    except requests.RequestException:
        return kalla, None, "Kunde inte nå källan"
    if not svar.ok:
        return kalla, None, f"Källan svarade med fel {svar.status_code}"
    feed = feedparser.parse(svar.content)
    if not feed.entries:
        return kalla, None, "Hittade inga artiklar, är det verkligen en RSS-länk?"
    return kalla, feed.entries, None


def _datum(entry):
    t = entry.get("published_parsed") or entry.get("updated_parsed")
    if not t:
        return None
    return datetime.fromtimestamp(calendar.timegm(t), tz=timezone.utc).astimezone(TIDSZON)


def hamta_nyheter(b):
    matcha = bygg_matchare(b)
    with ThreadPoolExecutor(max_workers=8) as pool:
        svar = list(pool.map(_hamta_en, b["kallor"]))

    resultat, statistik, sedda = [], [], set()
    for kalla, entries, fel in svar:
        if fel:
            statistik.append({"Källa": kalla["namn"], "Artiklar": 0, "Träffar": 0, "Status": fel})
            continue
        traffar = 0
        for e in entries:
            titel = _rensa_html(e.get("title", ""))
            sammanfattning = _rensa_html(e.get("summary", ""))
            kategorier = [t.get("term", "") for t in (e.get("tags") or [])]
            if kalla["filtrera"]:
                orsak, monster = matcha(titel, sammanfattning, kategorier)
                if not orsak:
                    continue
            else:
                orsak, monster = "hela källan", []
            lank = e.get("link", "")
            nyckel = lank or f"{kalla['namn']}:{titel}"
            if nyckel in sedda:
                continue
            sedda.add(nyckel)
            traffar += 1
            d = _datum(e)
            resultat.append({
                "Källa": kalla["namn"], "Titel": titel, "Länk": lank,
                "Publicerad": d.strftime("%Y-%m-%d %H:%M") if d else "Okänt datum",
                "Sammanfattning": sammanfattning or "Ingen sammanfattning",
                "Matchande nyckelord": orsak,
                "_tid": d.timestamp() if d else 0, "_monster": monster,
            })
        statistik.append({"Källa": kalla["namn"], "Artiklar": len(entries), "Träffar": traffar, "Status": "OK"})
    return pd.DataFrame(resultat), pd.DataFrame(statistik)


def markera(text, monster):
    """Gör om text till säker HTML där matchande ord är markerade."""
    if not monster:
        return html.escape(text)
    delar = re.split(f"({'|'.join(monster)})", text, flags=re.I)
    return "".join(f"<mark>{html.escape(d)}</mark>" if i % 2 else html.escape(d) for i, d in enumerate(delar))


# ---------- Excel ----------

def _export(df):
    return df.drop(columns=[c for c in df.columns if c.startswith("_")])


def skapa_excel(df):
    buf = BytesIO()
    _export(df).to_excel(buf, index=False)
    buf.seek(0)
    wb = load_workbook(buf)
    ws = wb.active
    for kol, bredd in zip("ABCDEF", [22, 60, 45, 18, 80, 26]):
        ws.column_dimensions[kol].width = bredd
    for rad in range(2, ws.max_row + 1):
        cell = ws.cell(row=rad, column=3)
        if cell.value:
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"
    ut = BytesIO()
    wb.save(ut)
    return ut.getvalue()


# ---------- Word ----------

def _bokmarke(paragraph, namn, id_):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(id_))
    start.set(qn("w:name"), namn)
    paragraph._p.insert(0, start)
    slut = OxmlElement("w:bookmarkEnd")
    slut.set(qn("w:id"), str(id_))
    paragraph._p.append(slut)


def _lank(paragraph, text, *, ankare=None, url=None):
    lank = OxmlElement("w:hyperlink")
    if ankare:
        lank.set(qn("w:anchor"), ankare)
    else:
        rid = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        lank.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    farg = OxmlElement("w:color")
    farg.set(qn("w:val"), "0563C1")
    rpr.append(farg)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    run.append(t)
    lank.append(run)
    paragraph._p.append(lank)


def skapa_word(df, statistik, rubrik):
    doc = Document()
    doc.add_heading(f"{rubrik}: {date.today():%Y-%m-%d}", level=1)
    kallor = [k for k in statistik.itertuples() if k.Träffar > 0]
    doc.add_paragraph(f"Totalt {antal(len(df), 'artikel', 'artiklar')} från {antal(len(kallor), 'källa', 'källor')}.")

    doc.add_heading("Översikt: antal träffar per källa", level=2)
    for i, k in enumerate(kallor):
        _lank(doc.add_paragraph(), f"{k.Källa}: {antal(k.Träffar, 'träff', 'träffar')}", ankare=f"kalla_{i}")

    for i, k in enumerate(kallor):
        h = doc.add_heading(f"{k.Källa} ({antal(k.Träffar, 'träff', 'träffar')})", level=2)
        _bokmarke(h, f"kalla_{i}", i + 1)
        for _, rad in df[df["Källa"] == k.Källa].sort_values("_tid", ascending=False).iterrows():
            doc.add_paragraph().add_run(rad["Titel"]).bold = True
            doc.add_paragraph(f"Publicerad: {rad['Publicerad']}")
            doc.add_paragraph(f"Ingress: {rad['Sammanfattning']}")
            p = doc.add_paragraph("Länk: ")
            if rad["Länk"]:
                _lank(p, rad["Länk"], url=rad["Länk"])
            doc.add_paragraph(f"Matchning: {rad['Matchande nyckelord']}")
            doc.add_paragraph("")

    buf = BytesIO()
    doc.save(buf)
    return buf.getvalue()
